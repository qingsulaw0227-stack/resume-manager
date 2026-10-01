import base64,copy,io,json,tempfile,unittest,sys,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import server as app
from lxml import etree as E
from docx import Document
class AppTest(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.old=(app.DATA,app.STATE);app.DATA=Path(self.tmp.name);app.STATE=app.DATA/'state.json';(app.DATA/'generated').mkdir()
 def tearDown(self):app.DATA,app.STATE=self.old;self.tmp.cleanup()
 def post(self,path,payload):
  h=app.Handler.__new__(app.Handler);h.path='/api/'+path;raw=json.dumps(payload).encode();h.headers={'Content-Length':str(len(raw))};h.rfile=io.BytesIO(raw);result=[];h.response=lambda obj,status=200:result.append((status,obj));h.do_POST();return result[0]
 def seed(self):
  self.assertEqual(self.post('person',{'person':{'name':'测试同学','english_name':'Test Student','contact':'student@example.com'}})[0],200)
  e={'title':'示例法律诊所｜志愿者','date':'2025.01–2025.02','group':'社会实践','raw':'完整的原始素材，不应直接输出。','description':'法律研究：协助整理资料。','english_title':'Example Legal Clinic | Volunteer','english_date':'Jan 2025–Feb 2025','english':'Legal research: Assisted with research.'}
  code,entry=self.post('entry',{'entry':e});self.assertEqual(code,200);return entry
 def test_empty_database(self):self.assertEqual(app.state()['entries'],[]);self.assertEqual(app.state()['person'],{})
 def test_scenarios_and_docx(self):
  e=self.seed();self.assertEqual(self.post('profile',{'name':'法律诊所申请'})[0],200)
  code,out=self.post('generate',{'profile':'法律诊所申请','ids':[e['id']],'language':'both','auto_layout':False});self.assertEqual(code,200,out)
  with zipfile.ZipFile(app.DATA/'generated'/out['name']) as z:
   r=E.fromstring(z.read('word/document.xml'));text=app.txt(r)
  self.assertIn('测试同学',text);self.assertIn('Example Legal Clinic',text);self.assertNotIn('原始素材',text)
  self.assertEqual(len(app.state()['history']),1)
  self.assertEqual(app.state()['history'][0]['entries_snapshot'][0]['raw'],e['raw'])
 def test_edit_requires_translation_review(self):
  e=self.seed();e['description']='法律研究：协助核对资料。';code,e=self.post('entry',{'entry':e});self.assertEqual(code,200)
  self.assertTrue(e['english_pending']);self.assertEqual(len(e['revisions']),1)
  code,_=self.post('generate',{'profile':'综合','ids':[e['id']],'language':'both','auto_layout':False});self.assertEqual(code,400)
 def test_upload_and_validation(self):
  code,out=self.post('import',{'name':'notes.txt','content':base64.b64encode('完整材料'.encode()).decode()});self.assertEqual(code,200);self.assertEqual(out['text'],'完整材料')
  self.assertEqual(self.post('profile',{'name':'../escape'})[0],400)
  self.assertEqual(self.post('import',{'name':'bad.docx','content':base64.b64encode(b'not a zip').decode()})[0],400)
 def test_no_selected_entries_no_leakage(self):
  e=self.seed();code,out=self.post('generate',{'profile':'综合','ids':[],'language':'zh','auto_layout':False});self.assertEqual(code,200)
  with zipfile.ZipFile(app.DATA/'generated'/out['name']) as z:self.assertNotIn('示例法律诊所',app.txt(E.fromstring(z.read('word/document.xml'))))
 def test_bulk_resume_import_deduplicates_and_flags_date_conflict(self):
  def resume(date,detail):
   doc=Document();doc.add_paragraph('实习经历');doc.add_paragraph('示例律师事务所｜实习生 '+date);doc.add_paragraph(detail)
   stream=io.BytesIO();doc.save(stream);return base64.b64encode(stream.getvalue()).decode()
  payload={'files':[
   {'name':'old.docx','content':resume('2024.01–2024.03','完成法律检索与合同审阅。')},
   {'name':'new.docx','content':resume('2024.01–2024.04','完成法律检索、合同审阅及备忘录起草。')},
  ]}
  code,out=self.post('bulk-import',payload);self.assertEqual(code,200,out)
  self.assertEqual(out['added'],1);self.assertEqual(len(out['entries']),1)
  imported=out['entries'][0]
  self.assertEqual(imported['group'],'实习工作');self.assertEqual(len(imported['sources']),2)
  self.assertTrue(imported['conflicts']);self.assertIn('备忘录',imported['description'])
 def test_long_title_uses_adaptive_columns_and_keeps_short_role_together(self):
  self.post('person',{'person':{'name':'测试同学','english_name':'Test Student','contact':'student@example.com'}})
  long_entry={'title':'全国大学生市场调查与分析大赛｜社区团购市场发展现状与可持续发展策略研究｜核心成员','date':'2022.11–2023.05','group':'科研课题','raw':'', 'description':'运用模型完成研究。','english_title':'Competition | Research Project | Core Member','english_date':'Nov 2022 – May 2023','english':'Conducted research.'}
  _,saved=self.post('entry',{'entry':long_entry});_,out=self.post('generate',{'profile':'综合','ids':[saved['id']],'language':'zh','auto_layout':False})
  with zipfile.ZipFile(app.DATA/'generated'/out['name']) as z:root=E.fromstring(z.read('word/document.xml'))
  xml=E.tostring(root,encoding='unicode');self.assertIn('核\u2060心\u2060成\u2060员',xml);self.assertIn('w:cantSplit',xml)
  grids=root.xpath('.//w:tblGrid/w:gridCol/@w:w',namespaces=app.NS)
  self.assertNotEqual(grids[-2:],['7938','2268'])
if __name__=='__main__':unittest.main()
