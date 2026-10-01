"""Editable DOCX from owner data; no personal template or fixed XML offsets."""
import io,re,base64,unicodedata
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
GROUPS={'教育背景':'EDUCATION','实习工作':'LEGAL EXPERIENCE','社会实践':'ACTIVITIES','科研课题':'RESEARCH','荣誉与技能':'HONORS & SKILLS','兴趣爱好':'INTERESTS'}
def font(run,size=10.5,bold=False,italic=False):
 run.font.name='Liberation Serif';run.font.size=Pt(size);run.bold=bold;run.italic=italic
 run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
def paragraph(p,space=0):
 p.paragraph_format.space_before=Pt(space);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(13)
 p.paragraph_format.widow_control=True
 return p
def table(doc,widths):
 t=doc.add_table(rows=1,cols=len(widths));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for cell,w in zip(t.rows[0].cells,widths):
  cell.width=Cm(w);pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
  for side in ['top','bottom','left','right']:
   el=OxmlElement('w:'+side);el.set(qn('w:w'),'0');el.set(qn('w:type'),'dxa');m.append(el)
  pr.append(m)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'nil');borders.append(el)
 t._tbl.tblPr.append(borders)
 return t
def heading(doc,text):
 p=paragraph(doc.add_paragraph(),4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
 r=p.add_run(text);font(r,14);r.font.color.rgb=RGBColor.from_string('1F4E79')
 borders=OxmlElement('w:pBdr');line=OxmlElement('w:bottom')
 for key,value in [('val','single'),('sz','5'),('space','2'),('color','C0504D')]:line.set(qn('w:'+key),value)
 borders.append(line);p._p.get_or_add_pPr().append(borders)
def description(p,text):
 # Only explicit short labels receive underline; original wording is retained.
 text=' '.join(text.splitlines());pos=0
 for m in re.finditer(r'(?:^|[。；;]\s*)([^。；;：:]{1,24}[：:])',text):
  a,b=m.span(1);font(p.add_run(text[pos:a]));r=p.add_run(text[a:b]);font(r);r.underline=True;pos=b
 font(p.add_run(text[pos:]))
def visual_units(text):
 """Approximate rendered width in ems; CJK glyphs count as full-width."""
 return sum(1 if unicodedata.east_asian_width(c) in 'WFA' else .56 if c.isalnum() else .45 for c in text)
def adaptive_columns(title,date,total=18):
 # Dates need only their measured width. The remainder belongs to the title.
 date_cm=min(4.4,max(2.35,visual_units(date)*10.5*0.0353+.35)) if date else 2.35
 return total-date_cm,date_cm
def title_size(title,title_cm):
 needed=visual_units(title)*10.5*0.0353
 return max(9.5,min(10.5,10.5*title_cm/max(needed,.1)))
def keep_short_role(text):
 """Keep a short final role such as 核心成员 together without visible spaces."""
 if 1<len(text)<=8 and all('\u4e00'<=c<='\u9fff' for c in text):return '\u2060'.join(text)
 return text
def add_title(p,title,lang,size):
 parts=[x.strip() for x in re.split('[｜|]',title)]
 for index,part in enumerate(parts):
  if index:font(p.add_run('｜' if lang=='zh' else ' | '),size,bold=lang=='zh')
  shown=keep_short_role(part) if index==len(parts)-1 else part
  font(p.add_run(shown),size,bold=(lang=='zh' or index==0),italic=(lang=='en' and index>0))
def generate_docx(person,entries,ids,language,path):
 if language not in ['zh','en','both']:raise ValueError('请选择有效的输出语言')
 if len(set(ids))!=len(ids):raise ValueError('经历选择重复')
 byid={e['id']:e for e in entries}
 if any(i not in byid for i in ids):raise ValueError('经历不存在')
 chosen=[byid[i] for i in ids]
 if not person.get('name'):raise ValueError('请先保存个人信息')
 if language in ['en','both'] and not person.get('english_name'):raise ValueError('请先填写英文姓名')
 for e in chosen:
  if not e.get('description'):raise ValueError('请填写“'+e['title']+'”的简历表述')
  if language in ['en','both'] and (e.get('english_pending') or not e.get('english') or not e.get('english_title') or (e.get('date') and not e.get('english_date'))):raise ValueError('“'+e['title']+'”英文缺失或待核对')
 doc=Document();doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='Resume'
 for ix,lang in enumerate(['zh','en'] if language=='both' else [language]):
  sec=doc.sections[0] if ix==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
  sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(1.5)
  sec.header_distance=sec.footer_distance=Cm(.5)
  photo=person.get('photo','');t=table(doc,[15.5,2.5] if photo else [18])
  p=paragraph(t.cell(0,0).paragraphs[0]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(30)
  r=p.add_run(person['name'] if lang=='zh' else person['english_name']);font(r,25);r.font.color.rgb=RGBColor.from_string('1F4E79')
  p=paragraph(t.cell(0,0).add_paragraph());p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run(person.get('contact','')),10)
  if photo:t.cell(0,1).paragraphs[0].add_run().add_picture(io.BytesIO(base64.b64decode(photo)),width=Cm(2))
  for group,en in GROUPS.items():
   group_entries=[e for e in chosen if e['group']==group]
   if not group_entries:continue
   heading(doc,group if lang=='zh' else en)
   for j,e in enumerate(group_entries):
    title=e['title'] if lang=='zh' else e['english_title'];date=e.get('date','') if lang=='zh' else e.get('english_date','')
    title_cm,date_cm=adaptive_columns(title,date);t=table(doc,[title_cm,date_cm]);p=paragraph(t.cell(0,0).paragraphs[0],0 if j==0 else 4);p.paragraph_format.keep_with_next=True
    trpr=t.rows[0]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
    add_title(p,title,lang,title_size(title,title_cm))
    p=paragraph(t.cell(0,1).paragraphs[0],0 if j==0 else 4);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.keep_with_next=True
    font(p.add_run(date),bold=True)
    p=paragraph(doc.add_paragraph());description(p,e['description'] if lang=='zh' else e['english'])
 doc.save(path)
